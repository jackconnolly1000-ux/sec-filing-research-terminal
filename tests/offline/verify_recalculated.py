"""Compare source formula caches with independently recalculated disposable LibreOffice copies."""
import argparse,json,math
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]

def verify(folder,examples=True):
    originals=[ROOT/'samples'/f'{s}_liquidity.xlsx' for s in ('lly','oracle','vistra')]
    if examples:originals += [ROOT/'outputs/final-examples'/n for n in ('programmatic-pilot.xlsx','research.xlsx','scenario.xlsx')]
    result=[]
    for original in originals:
        recalculated=folder/original.name
        if not recalculated.is_file():raise ValueError('Recalculated file is missing: '+str(recalculated))
        source=load_workbook(original,data_only=True);formulas=load_workbook(original,data_only=False);calc=load_workbook(recalculated,data_only=True)
        try:
            errors=[f'{s.title}!{c.coordinate}: {c.value}' for s in calc for row in s for c in row if c.data_type=='e']
            if errors:raise ValueError((original.name,errors))
            count=0
            for s in formulas:
                for row in s:
                    for cell in row:
                        if cell.data_type!='f':continue
                        before=source[s.title][cell.coordinate].value;after=calc[s.title][cell.coordinate].value
                        if isinstance(before,(int,float)):
                            if not isinstance(after,(int,float)) or not math.isclose(before,after,rel_tol=1e-10,abs_tol=1e-5):raise ValueError((original.name,s.title,cell.coordinate,before,after))
                        elif before!=after:raise ValueError((original.name,s.title,cell.coordinate,before,after))
                        count+=1
            result.append(dict(file=original.name,sheets=len(calc.sheetnames),formula_caches_verified=count,formula_errors=0))
        finally:source.close();formulas.close();calc.close()
    return result
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('folder',type=Path);p.add_argument('--companies-only',action='store_true');args=p.parse_args()
    print(json.dumps(verify(args.folder,not args.companies_only),indent=2))
