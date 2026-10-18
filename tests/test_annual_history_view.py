"""Annual presentation reconciles every app/export value to checked input versions."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import annual_history_view as a
import company_review
import oracle_liquidity
import vistra_liquidity

ROOT = Path(__file__).resolve().parents[1]


class AnnualPresentation(unittest.TestCase):
    def test_five_full_years_and_source_bases_are_visible(self):
        for ticker in ('LLY', 'ORCL', 'VST'):
            m=a.build(ticker);self.assertEqual(len(m['rows']),5)
            html=a.render(ticker)
            self.assertIn('USD millions',html);self.assertIn('TTM',html)
            for r in m['rows']:
                self.assertIn(r['source']['period_end'],html)
                self.assertIn(r['source']['accession'],html)
            self.assertIn('not',m['notes'][0])
        self.assertIn('classification, not new cash generation',a.render('LLY'))
        self.assertIn('reported zero',a.render('VST'))
        self.assertIn('(23,686.0)',a.render('ORCL'))

    def test_bad_years_calculations_and_periods_are_refused(self):
        original=json.loads(a.SAMPLE.read_text())
        for kind in ('year','calculation','period'):
            data=deepcopy(original);row=data['observations'][0]
            if kind=='year':row['fiscal_year']=2020
            if kind=='calculation':row['metrics']['free_cash_flow']['value']+=1
            if kind=='period':row['metrics']['operating_cash_flow']['inputs'][0]['period_start']='2021-02-01'
            with tempfile.TemporaryDirectory() as folder:
                path=Path(folder)/'history.json';path.write_text(json.dumps(data))
                with patch.object(a,'SAMPLE',path),patch.object(a,'CHECKED_SHA256',hashlib.sha256(path.read_bytes()).hexdigest()):
                    with self.assertRaises(ValueError):a.build('LLY')

    def test_changed_checkpoint_or_code_refuses_downloads(self):
        snapshots=json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']
        with patch.object(a,'CHECKED_SHA256','changed'):
            with self.assertRaises(ValueError):a.build('LLY')
        with patch.object(a,'fingerprint',return_value={'changed':True}):
            self.assertIsNone(company_review.saved_workbook(snapshots,ticker='LLY'))
            self.assertIsNone(oracle_liquidity.workbook_bytes(oracle_liquidity.for_snapshot(snapshots[1])))
            self.assertIsNone(vistra_liquidity.workbook_bytes(vistra_liquidity.for_snapshot(snapshots[2])))

    def test_all_seventy_five_cached_values_and_formulas_match_the_shared_model(self):
        for ticker,stem in [('LLY','lly'),('ORCL','oracle'),('VST','vistra')]:
            p=ROOT/f'samples/{stem}_liquidity.xlsx'
            values=load_workbook(p,data_only=True);formulas=load_workbook(p,data_only=False)
            try:
                for col,row in enumerate(a.build(ticker)['rows'],2):
                    for r,key in enumerate(a.KEYS,6):
                        self.assertEqual(values['Annual history'].cell(r,col).value,row['metrics'][key]['value'])
                        self.assertEqual(formulas['Annual history'].cell(r,col).data_type,'f')
                    letter='BCDEF'[col-2]
                    self.assertEqual(formulas['Annual history'][letter+'10'].value,f'={letter}8-{letter}9')
                self.assertEqual(len(formulas['Annual history'].merged_cells.ranges),0)
                manifest=json.loads(p.with_suffix('.manifest.json').read_text())
                self.assertEqual(manifest['annual_history_sha256'],a.fingerprint())
                data=values['Data'];first,last=manifest['annual_data_start']+1,manifest['annual_data_last']
                self.assertEqual(last-first+1,35 if ticker=='VST' else 25)
                urlcol=12 if ticker=='LLY' else 14 if ticker=='VST' else 13
                accessioncol=7 if ticker=='LLY' else 8
                for r in range(first,last+1):
                    self.assertIsInstance(data.cell(r,2).value,(int,float))
                    self.assertTrue(data.cell(r,urlcol).value.startswith('https://www.sec.gov/Archives/'))
                    self.assertTrue(data.cell(r,accessioncol).value)
                if ticker=='LLY':
                    self.assertNotIn('p. 56; cash flows p. 58',str([r for r in data.values]))
            finally:values.close();formulas.close()


class AnnualCharts(unittest.TestCase):
    """Two five-year charts per company, drawn from the same checked rows in the app and in Excel."""
    STEMS = {'LLY': 'lly', 'ORCL': 'oracle', 'VST': 'vistra'}

    def test_app_draws_both_charts_from_the_checked_rows(self):
        import html as h
        for ticker in self.STEMS:
            m = a.build(ticker); shown = h.unescape(a.charts(m))
            self.assertIn('Fiscal-year-end cash and baseline borrowing', shown)
            self.assertIn('Annual operating cash flow, capital spending and FCF', shown)
            for r in m['rows']:
                for key in ('cash', 'gross_debt', 'operating_cash_flow', 'free_cash_flow'):
                    self.assertIn(a._usd(r['metrics'][key]['value']), shown)
                self.assertIn(a._usd(-r['metrics']['capital_spending']['value']), shown)   # a use, left of zero
            self.assertEqual(shown.count('class="lr-bal-row"'), 10)
            self.assertEqual(shown.count('class="lr-bar-row"'), 15)
        self.assertIn('−$206M', h.unescape(a.charts(a.build('VST'))))           # FY2021 OCF stays negative
        # Charts follow the table and precede the reclassification and source notes
        full = h.unescape(a.render('LLY'))
        self.assertLess(full.index('</table>'), full.index('Fiscal-year-end cash and baseline borrowing'))
        self.assertLess(full.index('Annual operating cash flow'), full.index('classification, not new cash generation'))

    def test_a_missing_value_is_a_gap_not_a_zero_bar(self):
        m = deepcopy(a.build('VST'))
        m['rows'][1]['metrics']['gross_debt']['value'] = None
        m['rows'][2]['metrics']['free_cash_flow']['value'] = None
        shown = a.charts(m)
        self.assertEqual(shown.count('>Unavailable<'), 2)                     # visible values, one per gap
        import re
        track = re.search(r'title="Baseline borrowing, Dec 31, 2022: Unavailable">(.*?)</div>', shown).group(1)
        self.assertEqual(track, '')                                              # an empty track: no bar at all
        self.assertNotIn('$0', shown)

    def test_excel_charts_bind_to_the_annual_table_and_cache_its_values(self):
        import zipfile, re
        for ticker, stem in self.STEMS.items():
            path = ROOT / f'samples/{stem}_liquidity.xlsx'
            ws = load_workbook(path)['Annual history']; vs = load_workbook(path, data_only=True)['Annual history']
            self.assertEqual(len(ws._charts), 2)
            for c in 'BCDEF':
                self.assertEqual(ws[f'{c}61'].value, f'=IF(ISNUMBER({c}9),-{c}9,NA())')
                self.assertEqual(vs[f'{c}61'].value, -vs[f'{c}9'].value)
            m = a.build(ticker)
            with zipfile.ZipFile(path) as z:
                charts = [z.read(n).decode() for n in z.namelist() if n.startswith('xl/drawings/charts/') and "'Annual history'" in z.read(n).decode()]
            self.assertEqual(len(charts), 2)
            refs = sorted(set(re.findall(r"<c:f>'Annual history'!\$B\$(\d+):\$F\$\1</c:f>", ''.join(charts))))
            self.assertEqual(refs, sorted(['5', '6', '7', '8', '10', '61']))      # FY labels, the table rows, the helper row
            balances = next(x for x in charts if 'Fiscal-year-end cash' in x)
            cash = [float(v) for v in re.findall(r'<c:pt idx="\d"><c:v>([-\d.e+]+)</c:v>', balances.split('$B$6:$F$6')[1].split('</c:numCache>')[0])]
            self.assertEqual(cash, [r['metrics']['cash']['value'] for r in m['rows']])
            for x in charts:
                self.assertIn('<c:dispBlanksAs val="gap"/>', x)

    def test_chart_writer_leaves_a_missing_point_out(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location('charts', ROOT / 'tests/offline/add_annual_history_charts.py')
        mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
        labels = ['FY2021', 'FY2022', 'FY2023', 'FY2024', 'FY2025']
        values = {6: ('Cash', [1.0, None, 3.0, 4.0, 5.0]), 7: ('Borrowing', [1.0, 2.0, 3.0, 4.0, 5.0])}
        xml = mod.chart_xml('balances', labels, values)
        cash = xml.split('$B$6:$F$6')[1].split('</c:numCache>')[0]
        self.assertNotIn('<c:pt idx="1">', cash)                                # no cached point: no column
        self.assertIn('<c:ptCount val="5"/>', cash)
