"""Source-checked current funding disaggregation, with totals that cannot double count."""
from copy import deepcopy
import json
from pathlib import Path
import unittest
from openpyxl import load_workbook
import funding_detail as fd
from fact_registry import FactRegistry, fact_id
ROOT=Path(__file__).resolve().parents[1]

class FundingDetail(unittest.TestCase):
    def setUp(self):
        self.snapshots={s['ticker']:s for s in json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']}

    def test_current_funding_breakdowns_reconcile_without_increasing_debt(self):
        a=fd.build(self.snapshots['LLY']);self.assertEqual([r['value'] for r in a['rows']],[5284e6,1766e6,7050e6])
        self.assertEqual(a['rows'][0]['source']['context_id'],'c-295')
        self.assertEqual(len(a['rows'][0]['source']['dimensions']),3)
        self.assertNotEqual(a['rows'][0]['value'],5280e6)
        b=fd.build(self.snapshots['ORCL']);self.assertEqual(b['rows'][0]['value'],7625e6)
        self.assertIsNone(b['rows'][1]['value']);self.assertIsNone(b['rows'][2]['value'])
        c=fd.build(self.snapshots['VST']);self.assertEqual([r['value'] for r in c['rows']],[1876e6,0,300e6,2176e6])
        self.assertIn('Reported zero',c['rows'][1]['status'])

    def test_liability_classification_keeps_reported_and_derived_distinct(self):
        for ticker,values in [('LLY',[47741e6,60663e6]),('ORCL',[47514e6,188549e6]),('VST',[10557e6,26552e6])]:
            m=fd.build(self.snapshots[ticker]);self.assertEqual([f['value'] for f in m['liabilities']],values)
        f=fd.build(self.snapshots['VST'])['liabilities'][1]
        self.assertEqual(f['status'],'Derived');self.assertEqual(f['inputs'][0]['value']-f['inputs'][1]['value'],f['value'])
        s=deepcopy(self.snapshots['LLY']);s['metrics']['cash']['accession']='other'
        with self.assertRaises(ValueError):fd.build(s)

    def test_reported_zero_history_is_not_flattened(self):
        h=fd.build(self.snapshots['VST'])['history']
        self.assertEqual([r['short_term']['value']/1e6 for r in h],[0,650,0,0,1800])
        self.assertEqual([r['receivables']['value']/1e6 for r in h],[0,425,0,750,1225])

    def test_registry_distinguishes_dimensions_and_refuses_conflicting_values(self):
        f=fd.build(self.snapshots['LLY'])['rows'][0]['source'];reg=FactRegistry();key=reg.add(f)
        self.assertEqual(reg.add(deepcopy(f)),key)
        changed=dict(f,value=5280e6)
        with self.assertRaises(ValueError):reg.add(changed)
        changed['dimensions']=[];self.assertNotEqual(fact_id(changed),key)

    def test_workbook_funding_and_liability_blocks_equal_the_shared_model(self):
        for ticker,stem in [('LLY','lly'),('ORCL','oracle'),('VST','vistra')]:
            w=load_workbook(ROOT/f'samples/{stem}_liquidity.xlsx',data_only=True)
            try:
                m=fd.build(self.snapshots[ticker]);s=w['Borrowing and leases']
                for i,r in enumerate(m['rows'],29):self.assertEqual(s.cell(i,2).value,'Unavailable' if r['value'] is None else r['value'])
                self.assertEqual([s[f'B{r}'].value for r in (45,46)],[f['value'] for f in m['liabilities']])
                if ticker=='VST':self.assertEqual(s['B8'].value,0);self.assertIn('Reported zero',s['E8'].value);self.assertIn('\"0\"',s['B8'].number_format)
            finally:w.close()
