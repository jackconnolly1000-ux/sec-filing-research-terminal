"""All views reconcile to the same known principal, without CP or duplicated term-loan aggregate."""
import json
from pathlib import Path
import unittest
from copy import deepcopy
from openpyxl import load_workbook
import debt_detail as d
import oracle_principal as op
import borrowing_leases as b
import annual_history_view as a
ROOT=Path(__file__).resolve().parents[1]
class DebtDetail(unittest.TestCase):
    def test_matrix_composition_and_calendar_are_one_reconciliation(self):
        m=d.build();self.assertEqual(m['total'],125567e6)
        self.assertEqual(sum(sum(row) for row in m['matrix']),m['total'])
        self.assertEqual([sum(col) for col in zip(*m['matrix'])],[y['principal'] for y in m['years']])
        self.assertEqual(m['composition'][2]['principal_usd'],5067e6)
        self.assertEqual(sum(e['principal_usd'] for e in m['events'] if e['kind']=='Term loan'),5067e6)
        self.assertTrue(all(e['date']>op.ANCHOR for e in m['events']))
        self.assertNotIn('Commercial paper',m['kinds'])

    def test_coupon_bins_exclude_all_floating_spreads_and_amortization_rates(self):
        m=d.build();self.assertEqual(sum(x['principal_usd'] for x in m['bins']),m['composition'][0]['principal_usd'])
        self.assertEqual(sum(x['tranche_count'] for x in m['bins']),len(m['fixed']))
        self.assertTrue(all(r['kind']=='Fixed-rate senior notes' for r in m['fixed']))
        for bucket in m['bins']:
            if bucket['principal_usd']:self.assertTrue(bucket['coupon_low_pct']<=bucket['weighted_coupon_pct']<bucket['coupon_high_pct'])
            else:self.assertIsNone(bucket['weighted_coupon_pct'])
        r=deepcopy(op.load());next(x for x in r['remaining'] if x['kind']=='Fixed-rate senior notes')['stated_rate']=None
        with self.assertRaises(ValueError):d.build(r)

    def test_plotly_views_use_dollars_in_billions_and_document_exclusions(self):
        m=d.build()
        for name,kind in [('Maturity concentration','heatmap'),('Fixed-rate coupon distribution','bar'),('Instrument composition','pie')]:
            f=d.figure(m,name);self.assertEqual(f.data[0].type,kind)
        self.assertAlmostEqual(sum(d.figure(m,'Instrument composition').data[0].values),125.567)
        self.assertIn('Unavailable',m['note'])

    def test_total_liabilities_are_a_separate_stack_with_derived_vistra(self):
        for s in json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']:
            f,note=b.liability_figure(s);self.assertEqual(f.layout.barmode,'stack');self.assertEqual(len(f.data),2)
            self.assertEqual([trace.orientation for trace in f.data],['h','h'])
            self.assertEqual([trace.y[0] for trace in f.data],['Total liabilities']*2)
            if s['ticker']=='VST':self.assertAlmostEqual(f.data[1].x[0],26.552);self.assertIn('derived',note)

    def test_maturity_heatmap_aggregates_every_annual_fact_and_keeps_labels_visible(self):
        m=d.build();periods=d.maturity_periods(m);f=d.figure(m,'Maturity concentration')
        self.assertLessEqual(len(periods),10)
        self.assertAlmostEqual(sum(sum(row) for row in f.data[0].z),m['total']/1e9)
        self.assertGreaterEqual(f.layout.margin.l,150)
        self.assertNotEqual(f.layout.plot_bgcolor,'#ffffff')
        self.assertEqual([p['label'] for p in periods[:5]],['FY27 rem.','FY28','FY29','FY30','FY31'])

    def test_annual_flow_ratio_uses_same_fiscal_year_and_preserves_gaps(self):
        m=a.build('VST');f,note=a.funding_history(m)
        expected=[r['metrics']['operating_cash_flow']['value']/r['metrics']['gross_debt']['value'] for r in m['rows']]
        self.assertEqual(list(f.data[1].y),expected);self.assertLess(f.data[1].y[0],0)
        m['rows'][1]['metrics']['gross_debt']['value']=None
        self.assertIsNone(a.funding_history(m)[0].data[1].y[1]);self.assertIn('not leverage',note)

    def test_excel_bins_composition_and_heatmap_match_the_shared_model(self):
        m=d.build();w=load_workbook(ROOT/'samples/oracle_liquidity.xlsx',data_only=True)
        try:
            s=w['Debt detail'];self.assertEqual([s.cell(r,2).value for r in range(32,35)],[c['principal_usd'] for c in m['composition']])
            self.assertEqual([s.cell(r,2).value for r in range(15,26)],[c['principal_usd'] for c in m['bins']])
            for i,row in enumerate(m['matrix'],6):self.assertEqual([s.cell(i,k).value for k in range(2,42)],row)
            self.assertEqual(len(s._charts),2)
        finally:w.close()
