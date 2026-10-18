"""Charts must show the checked calendar and the selected run, never missing values as zero."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest
import numpy as np
from openpyxl import load_workbook
import liquidity_simulation as ls
import oracle_principal as op
import oracle_principal_view as ov
import company_review
import scenario_export as se
import simulation_view as sv

ROOT=Path(__file__).resolve().parents[1]


class OracleCharts(unittest.TestCase):
    def test_every_fiscal_year_is_one_column_on_one_scale(self):
        # Replaces the cut "FY2032 onward" bar: each year from the FY2027 remainder to FY2066 has its own column
        review=op.load();years=ov.ladder_years(review)
        self.assertEqual([y['fy'] for y in years],list(range(2027,2067)))
        self.assertEqual(sum(y['principal'] for y in years),review['reconciliation']['known_principal_excluding_cp'])
        by_bucket={b['label']:b['principal'] for b in review['calendar']}
        self.assertEqual(years[0]['principal'],by_bucket['Remainder FY2027'])
        self.assertEqual(years[1]['principal'],by_bucket['FY2028'])
        self.assertEqual(sum(y['principal'] for y in years if y['fy']>=2032),by_bucket['Thereafter'])
        self.assertEqual(max(y['principal'] for y in years),10_750_000_000)       # one scale, nothing to cut
        fig,caption=ov.ladder_figure(review)
        self.assertEqual(list(fig.data[0].y),[y['principal']/1e9 for y in years[:10]])
        self.assertEqual(list(fig.layout.xaxis.tickvals),[y['label'] for y in years[:10]])
        self.assertEqual(fig.layout.xaxis.ticktext[0],'FY27*')
        self.assertEqual(fig.layout.xaxis.tickangle,-45)
        self.assertIn('complete calendar, including later years',caption)
        self.assertIn('$7.317B due September 2026–August 2027 is the FY2027 remainder ($2.672B) plus June–August 2027 ($4.645B',caption)
        self.assertIn('commercial paper (Unavailable)',caption)
        bad=deepcopy(review);bad['remaining'][3]['principal_aug_31']+=1
        with self.assertRaises(ValueError):ov.ladder_years(bad)                    # must add up to the checked total

    def test_cash_bridges_reconcile_list_additions_first_and_never_go_below_zero(self):
        import json,lly_liquidity_view as lv,oracle_liquidity_view as olv
        snaps={x['ticker']:x for x in json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']}
        for fig,first,last in ((olv.bridge_figure(snaps['ORCL'])[0],'$31.289B','$38.934B'),
                               (lv.bridge_figure(snaps['LLY'])[0],'$7.268B','$8.950B')):
            w=fig.data[0];text=list(w.text);x=list(w.x)
            self.assertEqual(w.type,'waterfall')
            self.assertEqual((w.customdata[0],w.customdata[-1]),(first,last))
            self.assertEqual((text[0],text[-1]),(f'${x[0]:,.2f}B',f'${x[-1]:,.2f}B'))
            self.assertEqual(w.hovertemplate,'%{y}: %{customdata}<extra></extra>')
            self.assertEqual(list(w.measure),['absolute']+['relative']*(len(x)-2)+['total'])
            steps=x[1:-1];signs=[v>0 for v in steps]
            self.assertEqual(signs,sorted(signs,reverse=True))                       # additions, then uses
            running=x[0]
            for v in steps:
                running+=v;self.assertGreaterEqual(running,0)
            self.assertAlmostEqual(running,x[-1],places=6)
            import theme as T
            self.assertEqual((w.increasing.marker.color,w.decreasing.marker.color),(T.FLOW_IN,T.FLOW_OUT))   # direction, not good/bad
        with self.assertRaises(ValueError):
            olv.waterfall(1e9,[("Use",-2e9)],-1e9,"Open","Close")                   # negative in any order: refused

    def test_thereafter_row_is_set_apart_as_combined_years(self):
        import json,html,oracle_liquidity_view as olv
        snap=next(x for x in json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots'] if x['ticker']=='ORCL')
        text=html.unescape(olv.repayment_html(snap,company_review.money))
        head=text.index('class="ol-later-head"');row=text.index('class="ol-later"')
        self.assertLess(text.index('FY2032'),head);self.assertLess(head,row)
        self.assertIn('Later years, combined',text);self.assertIn('Thereafter (combined)',text)
        self.assertIn('From Jun 2032, all later years',text)

    def test_histogram_splits_at_zero_and_counts_every_trial(self):
        i=ls.oracle_inputs();r=ls.simulate(i,ls.oracle_decisions()[1],5000,7)
        fig,size=sv.histogram_figure(r)
        below,above=fig.data
        self.assertEqual(len(below.x)+len(above.x),5000)
        self.assertTrue((np.asarray(below.x)<0).all() and (np.asarray(above.x)>=0).all())
        self.assertAlmostEqual(round(-below.xbins.start/size,9)%1,0)    # zero is a bin edge, so no bar straddles it
        self.assertEqual(fig.layout.yaxis.title.text,'Simulated trials (of 5,000)')
        self.assertEqual(fig.layout.xaxis.title.text,'Ending cash · USD billions')
        self.assertIn(f'{r["share_below_zero"]:.1%}',below.name)

    def test_monthly_band_is_independently_computed_from_the_trials(self):
        i=ls.oracle_inputs();d=ls.oracle_decisions()[1];p=ls.monthly_path(i,d,1000,42)
        for q in (10,50,90):np.testing.assert_array_equal(p[f'month_p{q}'],np.percentile(p['path'],q,axis=0))
        f=sv.monthly_figure(p)
        np.testing.assert_array_equal(f.data[2].y,p['month_p50']/ls.B)
        self.assertEqual(list(f.data[0].x),list(ls.MONTH_ENDS))
        self.assertEqual(f.data[1].fill,'tonexty')
        self.assertEqual(f.layout.shapes[0].y0,0)
        # Month labels name the month of each month-end point, never the day ("Oct 26") or the next month
        self.assertEqual(list(f.layout.xaxis.ticktext),['Sep 2026','Jan 2027','May 2027','Aug 2027'])
        self.assertEqual(list(f.layout.xaxis.tickvals),['2026-09-30','2027-01-31','2027-05-31','2027-08-31'])
        self.assertTrue(all(list(t.customdata)==list(ls.MONTHS) for t in f.data))

    def test_scenario_chart_follows_the_selected_decision_and_seed(self):
        i=ls.oracle_inputs();ds=ls.oracle_decisions();rs=[(ls.simulate(i,d,1000,42),d) for d in ds]
        monthly=[(d,r,ls.monthly_path(i,d,1000,42)) for r,d in rs]
        raw=se.workbook(i,rs,1000,42,ls.compare_correlations(i,ds[2],1000,42),monthly=monthly)
        w=load_workbook(BytesIO(raw));chart=w['Monthly path by month']._charts[0]
        self.assertIn('$F$26:$F$37',chart.series[0].val.numRef.f)
        for series,field in zip(chart.series[:3],('month_p10','month_p50','month_p90')):
            np.testing.assert_allclose([float(p.v) for p in series.val.numRef.numCache.pt],monthly[2][2][field])
        w.close()

    def test_saved_workbooks_state_dimensions_and_units_in_words(self):
        for name in ('lly_liquidity','oracle_liquidity','vistra_liquidity'):
            w=load_workbook(ROOT/f'samples/{name}.xlsx',read_only=True)
            texts=[v for ws in w for row in ws.iter_rows(values_only=True) for v in row if isinstance(v,str)]
            w.close()
            self.assertNotIn('[]',texts,name)
            self.assertFalse([t for t in texts if t.startswith('[[')],name)          # no raw JSON dimension lists
            self.assertFalse([t for t in texts if 'unless x' in t],name)
            self.assertIn('None (consolidated total)',texts,name)
            self.assertTrue(any(t.startswith('USD millions, except ratios (x) and percentages (%).') for t in texts),name)

    def test_saved_oracle_ladder_is_scaled_like_the_app(self):
        w=load_workbook(ROOT/'samples/oracle_liquidity.xlsx');chart=w['Payment dates']._charts[0]
        import oracle_principal_view as view, oracle_principal as principal
        years=view.ladder_years(principal.load())
        self.assertEqual(chart.type,'col')
        self.assertEqual(len(chart.series[0].val.numRef.numCache.pt),len(years))
        self.assertEqual([float(p.v) for p in chart.series[0].val.numRef.numCache.pt],[y['principal'] for y in years])
        self.assertLess(max(y['principal'] for y in years),12e9)
        w.close()

    def test_saved_oracle_ladder_links_to_the_existing_formula_table(self):
        values=load_workbook(ROOT/'samples/oracle_liquidity.xlsx',data_only=True)
        chart=values['Payment dates']._charts[0]
        self.assertEqual(chart.series[0].val.numRef.f,"'Payment dates'!$I$62:$I$101")
        cached=[float(p.v) for p in chart.series[0].val.numRef.numCache.pt]
        np.testing.assert_allclose(cached,[values['Payment dates'].cell(r,9).value for r in range(62,102)])
        self.assertEqual(len(values['Core charts']._charts),2)
        values.close()
