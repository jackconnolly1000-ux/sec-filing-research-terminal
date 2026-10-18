import json
import unittest
from pathlib import Path
from openpyxl import load_workbook
import core_charts as c
import company_review
ROOT=Path(__file__).resolve().parents[1]


class CoreCharts(unittest.TestCase):
    def setUp(self):
        self.snapshots=json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']

    def test_signed_flows_and_missing_history_are_not_zero(self):
        for s in self.snapshots:
            model=c.build(s);m=s['metrics']
            self.assertEqual(model['flow'],[('Operating cash flow',m['operating_cash_flow']['value']),('Capital spending',-m['capital_spending']['value']),('Free cash flow',m['free_cash_flow']['value'])])
            self.assertEqual(model['series'],['Cash and equivalents','Baseline borrowing'])
            if s['ticker']=='VST':self.assertEqual(model['points'][0][1],[785e6,20068e6])

    def test_app_has_zero_basis_dates_and_no_double_counted_flow_instruction(self):
        for s in self.snapshots:
            import annual_history_view as annual
            text=company_review.ttm_equation(c.build(s)['flow'],company_review.money)+annual.render(s['ticker'])
            self.assertIn('FCF',text);self.assertIn('USD',text)
            figures=annual.figures(annual.build(s['ticker']),annual.latest_balances(s))
            self.assertEqual(figures[1].layout.shapes[0].y0,0)
            from datetime import date
            end=date.fromisoformat(s['metrics']['cash']['period_end'])
            self.assertIn(f"{end:%b} {end.day}, {end:%Y}",annual.chart_captions(annual.build(s["ticker"]),annual.latest_balances(s))[0])   # the app shows Jun 30, 2026; Excel keeps ISO dates

    def test_two_native_charts_and_source_formula_bindings(self):
        for s in self.snapshots:
            stem={'LLY':'lly','ORCL':'oracle','VST':'vistra'}[s['ticker']];p=ROOT/f'samples/{stem}_liquidity.xlsx';v=load_workbook(p,data_only=True);f=load_workbook(p,data_only=False)
            self.assertEqual(len(f['Core charts']._charts),2)
            for cell,ref in [('B6',"='Debt and liquidity'!B10"),('B7',"=-'Debt and liquidity'!B11"),('B8',"='Debt and liquidity'!B12"),('B25',"='History'!B6"),('B26',"='History'!C6")]:self.assertEqual(f['Core charts'][cell].value,ref)
            self.assertEqual(v['Core charts']['B7'].value,-s['metrics']['capital_spending']['value'])
            if s['ticker']=='VST':self.assertEqual(v['Core charts']['C25'].value,20068e6);self.assertEqual(len(f['Core charts']._charts[1].series),2)
            self.assertEqual(json.loads(p.with_suffix('.manifest.json').read_text())['charts_sha256'],c.fingerprint());v.close();f.close()
