"""Experimental research exports preserve the archived record and expose calculation inputs."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import unittest
from unittest import mock

from openpyxl import load_workbook
import research_export as record

ROOT = Path(__file__).resolve().parents[1]


class ResearchExport(unittest.TestCase):
    def setUp(self):
        self.profiles = [json.loads((ROOT / f"samples/{t}.json").read_text()) for t in ("AAPL", "STRL")]

    def test_saved_figures_and_catalysts_are_preserved_with_formulas_and_source_inputs(self):
        raw = record.workbook(self.profiles, "2026-10-17T18:00:00-04:00")
        formulas = load_workbook(BytesIO(raw))
        values = load_workbook(BytesIO(raw), data_only=True)
        self.addCleanup(formulas.close)
        self.addCleanup(values.close)
        self.assertEqual(values.sheetnames, ["Research record", "Market data", "Findings", "8-K events",
                                            "Evidence", "Filing review", "Data", "About"])
        self.assertEqual(values["Research record"]["B3"].value, "2026-10-17T18:00:00-04:00")
        for p in self.profiles:
            for metric in p["sec_metrics"]["metrics"].values():
                r = next(i for i in range(2, values["Market data"].max_row+1)
                         if values["Market data"].cell(i,1).value == p["ticker"]
                         and values["Market data"].cell(i,2).value == metric["label"]
                         and "SEC XBRL" in values["Market data"].cell(i,6).value)
                self.assertAlmostEqual(values["Market data"].cell(r,3).value, metric["value"])
                self.assertIn("'Data'!E", formulas["Market data"].cell(r,3).value)
                self.assertEqual(formulas["Market data"].cell(r,3).font.color.rgb, "FF008000")
            headlines = [i["headline"] for kind in ("catalysts","risks","other_risks")
                         for i in p["analysis"]["sections"][kind]["items"]]
            exported = [r[2] for r in values["Findings"].values if r[0] == p["ticker"]]
            self.assertEqual(headlines, exported)
        self.assertTrue(all(isinstance(r[5], datetime) for r in list(values["Data"].values)[1:]))
        for sheet in formulas:
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertFalse(sheet.merged_cells)
            self.assertEqual(sheet["A1"].font.name, "Arial")
        self.assertEqual(formulas["Data"].freeze_panes,"A2")
        self.assertEqual(formulas["Data"]["K2"].hyperlink.target, values["Data"]["K2"].value)

    def test_missing_values_stay_blank_and_reported_zero_is_preserved(self):
        p = deepcopy(self.profiles[0])
        p["financials"]["metrics"]["forward_pe"]["value"] = None
        p["financials"]["metrics"]["market_cap"]["value"] = 0
        w = load_workbook(BytesIO(record.workbook([p])), data_only=True)
        self.addCleanup(w.close)
        data = {r[1]:r for r in w["Market data"].values if r[0]=="AAPL" and "Yahoo" in str(r[5])}
        self.assertIsNone(data["Forward P/E"][2])
        self.assertEqual(data["Market cap"][2],0)
        self.assertEqual(data["Market cap"][4],"Period not stated")

    def test_inconsistent_component_cannot_be_exported_as_a_matching_formula(self):
        self.profiles[0]["sec_metrics"]["metrics"]["revenue"]["components"][0]["value"] += 1e9
        with self.assertRaisesRegex(ValueError,"does not reconcile"):
            record.workbook(self.profiles)

    def test_legacy_records_without_components_keep_an_explicit_limit(self):
        m = self.profiles[0]["sec_metrics"]["metrics"]["revenue"]
        m.pop("components")
        rows = record.tables(self.profiles)["Market data"]
        row = next(r for r in rows if r[:2]==["AAPL","Revenue"])
        self.assertEqual(row[2],m["value"])
        self.assertIn("component formula unavailable",row[5])

    def test_narrative_is_literal_text_and_changed_template_is_refused(self):
        self.profiles[0]["analysis"]["sections"]["catalysts"]["items"][0]["headline"]='=1+1'
        w=load_workbook(BytesIO(record.workbook(self.profiles)))
        self.addCleanup(w.close)
        self.assertEqual(w["Findings"]["C2"].value,'=1+1')
        self.assertEqual(w["Findings"]["C2"].data_type,'s')
        with mock.patch.object(Path,"read_bytes",return_value=b"changed"):
            with self.assertRaisesRegex(ValueError,"layout changed"):
                record.workbook(self.profiles)
